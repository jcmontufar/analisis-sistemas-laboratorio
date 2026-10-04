from pathlib import Path
import csv
import xml.etree.ElementTree as ET

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw
from pypdf import PdfReader
from reportlab.pdfgen import canvas

RAIZ = Path(__file__).resolve().parents[2]
SALIDA = RAIZ / "pruebas" / "salida"


def main():
    SALIDA.mkdir(parents=True, exist_ok=True)

    md = SALIDA / "formato.md"
    md.write_text("# Prueba de Markdown\n\nArchivo recuperable.\n", encoding="utf-8")

    csv_path = SALIDA / "formato.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as archivo:
        csv.writer(archivo).writerows([["campo", "valor"], ["estado", "válido"]])

    docx_path = SALIDA / "formato.docx"
    doc = Document()
    doc.add_heading("Prueba de DOCX", level=1)
    doc.add_paragraph("Archivo recuperable.")
    doc.save(docx_path)

    xlsx_path = SALIDA / "formato.xlsx"
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Prueba"
    hoja.append(["campo", "valor"])
    hoja.append(["estado", "válido"])
    libro.save(xlsx_path)

    pdf_path = SALIDA / "formato.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(72, 760, "Prueba de PDF recuperable")
    pdf.save()

    png_path = SALIDA / "formato.png"
    imagen = Image.new("RGB", (480, 160), "white")
    ImageDraw.Draw(imagen).text((30, 60), "Prueba de PNG", fill="black")
    imagen.save(png_path)

    svg_path = SALIDA / "formato.svg"
    svg_path.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="480" height="160"><text x="30" y="80">Prueba de SVG</text></svg>', encoding="utf-8")

    assert md.read_text(encoding="utf-8").startswith("# Prueba")
    assert list(csv.reader(csv_path.open(encoding="utf-8")))[1][1] == "válido"
    assert Document(docx_path).paragraphs[0].text == "Prueba de DOCX"
    assert load_workbook(xlsx_path)["Prueba"]["B2"].value == "válido"
    assert len(PdfReader(pdf_path).pages) == 1
    Image.open(png_path).verify()
    assert ET.parse(svg_path).getroot().tag.endswith("svg")
    print("Formatos comprobados: Markdown, CSV, DOCX, XLSX, PDF, PNG y SVG.")


if __name__ == "__main__":
    main()
