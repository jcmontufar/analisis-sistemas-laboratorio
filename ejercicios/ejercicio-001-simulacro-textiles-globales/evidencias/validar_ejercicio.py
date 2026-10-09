"""Verificaciones reproducibles; no certifican corrección académica por sí solas."""
from pathlib import Path
import xml.etree.ElementTree as E
import hashlib, json, re
import zipfile
from pypdf import PdfReader
from PIL import Image
from docx import Document

B=Path(__file__).resolve().parents[1]
checks=[]
def check(name,condition):
    checks.append({'comprobacion':name,'resultado':'PASS' if condition else 'FAIL'})
def cells(stem):
    root=E.parse(B/'diagramas/fuentes'/f'{stem}.drawio').getroot()
    return root,{c.get('id'):c for c in root.findall('.//mxCell')}
def norm(s):return ' '.join(s.split())
stems=['01-stack-tecnologias','02-dfd-contexto','03-dfd-nivel-0','04-casos-uso','05-secuencia','06-colaboracion']
for stem in stems:
    root,cs=cells(stem)
    model=root.find('.//mxGraphModel')
    w,h=float(model.get('pageWidth')),float(model.get('pageHeight'))
    check(stem+': DRAW.IO nativo editable',root.tag=='mxfile' and bool(cs))
    check(stem+': referencias de conectores',all(c.get('source') in cs and c.get('target') in cs for c in cs.values() if c.get('edge')=='1'))
    inside=True
    for c in cs.values():
        if c.get('vertex')!='1':continue
        g=c.find('mxGeometry'); x,y,ww,hh=[float(g.get(k,'0')) for k in ('x','y','width','height')]
        inside &= x>=0 and y>=0 and x+ww<=w and y+hh<=h
    check(stem+': vértices dentro del lienzo',inside)
    for ext in ('pdf','png','svg'):
        p=B/'diagramas/exportados'/f'{stem}.{ext}'
        check(stem+': exportación '+ext,p.exists() and p.stat().st_size>1000)
        if ext=='pdf':check(stem+': una página PDF',len(PdfReader(p).pages)==1)
        if ext=='png':Image.open(p).verify()
        if ext=='svg':E.parse(p)
ctx,cc=cells(stems[1]); zero,zc=cells(stems[2])
def flows(cs):return {c.get('flow_id'):(c.get('actor'),c.get('direction'),norm(c.get('value'))) for c in cs.values() if c.get('flow_id')}
check('Balance DFD F01–F08: actor, dirección y nombre',flows(cc)==flows(zc) and len(flows(cc))==8)
check('Contexto: un proceso y cuatro entidades',sum(c.get('kind')=='process' for c in cc.values())==1 and sum(c.get('kind')=='entity' for c in cc.values())==4)
check('Nivel 0: cuatro procesos y cuatro almacenes',sum(c.get('kind')=='process' for c in zc.values())==4 and sum(c.get('kind')=='store' for c in zc.values())==4)
edges=[c for c in zc.values() if c.get('edge')=='1']
check('DFD: sin entidad→almacén ni almacén→almacén',all('process' in (zc[c.get('source')].get('kind'),zc[c.get('target')].get('kind')) for c in edges))
check('DFD: procesos y almacenes con entrada y salida',all(any(e.get('target')==id for e in edges) and any(e.get('source')==id for e in edges) for id,c in zc.items() if c.get('kind') in ('process','store')))
_,uc=cells(stems[3])
inc=[c for c in uc.values() if c.get('relation')=='include']; ext=[c for c in uc.values() if c.get('relation')=='extend']
check('UML: cinco include hacia comportamiento incluido',len(inc)==5 and all(c.get('source').startswith('uc') and c.get('target').startswith('inc') for c in inc))
check('UML: extend rechazo hacia inspección con guarda',len(ext)==1 and ext[0].get('source')=='reject' and ext[0].get('target')=='uc3' and '[prenda rechazada]' in ext[0].get('value'))
_,sq=cells(stems[4]); _,co=cells(stems[5])
def messages(cs):return {int(c.get('message_id')):(c.get('sender'),c.get('receiver'),c.get('message')) for c in cs.values() if c.get('message_id')}
check('Secuencia/colaboración: 20 mensajes idénticos',messages(sq)==messages(co) and set(messages(sq))==set(range(1,21)))
# Controles puntuales de revisión 1.1, además de las validaciones originales.
model=json.loads((B/'evidencias/modelo.json').read_text(encoding='utf-8'))
for fid,label in [('F04','Datos de corte y estado del registro'),('F06','Datos de paquete habilitado y estado de confección')]:
    check('Revisión 1.1: nombre de '+fid,flows(cc)[fid][2]==fid+' '+label)
    check('Revisión 1.1: '+fid+' en análisis/matriz',all(label in (B/p).read_text(encoding='utf-8') for p in ['analisis/ANALISIS.md','validacion/CONSISTENCIA.md']))
check('Revisión 1.1: D4→P1 correctamente orientado',any(c.get('source')=='d4alias' and c.get('target')=='p0' and c.get('logical_source')=='db3' for c in edges) and zc['d4alias'].get('logical_store')=='db3')
check('Revisión 1.1: sin transferencia P4→P1',not any(c.get('source')=='p3' and c.get('target')=='p0' for c in edges))
check('Revisión 1.1: cuatro almacenes lógicos, alias explícito',sum(c.get('kind')=='store_alias' for c in zc.values())==1 and 'mismo D4' in zc['d4alias'].get('value'))
check('Revisión 1.1: liberación sin atribuir empaque físico',uc['uc4'].get('value')=='Liberar aprobadas para empaque' and 'no se le atribuye la ejecución del empaque físico' in (B/'analisis/ANALISIS.md').read_text(encoding='utf-8'))
check('Revisión 1.1: mensaje 20 confirma liberación',messages(sq)[20]==('4','3','Código de lote y confirmación de liberación'))
expected_old=['OT y metraje asignado','Guardar OT y habilitación','OT persistida','OT habilitada','Conteo de piezas y merma','Guardar corte y paquete','Corte persistido','Confirmación de corte','Avance por paquete y tiempo de lote','Guardar avance y tiempo','Confección persistida','Confirmación de confección','Resultado individual de inspección','[aprobada] Guardar prenda aprobada','[rechazada] Guardar rechazo y causa','Inspección persistida','Resumen de aprobadas y rechazadas','Liberación de aprobadas para empaque','[hay aprobadas] Generar y guardar código de lote']
check('Revisión 1.1: contenidos 1–19 conservados',all(messages(sq)[i][2]==label for i,label in enumerate(expected_old,1)))
check('Revisión 1.1: exactamente 20 mensajes sin duplicación',all(sum(bool(c.get('message_id')) for c in cs.values())==20 for cs in [sq,co]))
expected_pairs=[(0,4),(4,5),(5,4),(4,0),(1,4),(4,5),(5,4),(4,1),(2,4),(4,5),(5,4),(4,2),(3,4),(4,5),(4,5),(5,4),(4,3),(3,4),(4,5),(4,3)]
check('Revisión 1.1: emisores/receptores originales preservados',all(messages(sq)[i][:2]==tuple(map(str,pair)) for i,pair in enumerate(expected_pairs,1)))
check('Revisión 1.1: loop y alt conservados',sq['loop'].get('value')=='loop [por cada prenda]' and sq['alt'].get('value')=='alt' and messages(sq)[19][2].startswith('[hay aprobadas]'))
expected_flows={'F01':'Orden de Trabajo y datos de tela asignada','F02':'OT habilitada y resumen de producción','F03':'Registro de corte y merma','F05':'Registro de avance y tiempo por lote','F07':'Resultado de inspección y defectos','F08':'Resumen de calidad y código de lote'}
check('Revisión 1.1: otros seis flujos no alterados',all(flows(cc)[fid][2]==fid+' '+label for fid,label in expected_flows.items()))
check('Revisión 1.1: cuatro actores originales preservados',all([cs['a'+str(i)].get('value') for i in range(4)]==model['actors']==['Jefe de Planta','Operario de Corte','Operario de Confección','Inspector de Calidad'] for cs in [cc,zc,uc,co]))
links=[c for c in co.values() if c.get('link_messages')]
linked=[int(n) for c in links for n in c.get('link_messages').split(',')]
participant_ids={'a0':'0','a1':'1','a2':'2','a3':'3','system':'4','db':'5'}
check('Revisión 1.1: enlaces contienen 1–20 exactamente una vez',sorted(linked)==list(range(1,21)))
check('Revisión 1.1: llamadas numeradas y direcciones equivalentes',all(all(re.search(r'\b'+n+r':',c.get('value')) and messages(sq)[int(n)][:2]==(participant_ids[c.get('source')],participant_ids[c.get('target')]) for n in c.get('link_messages').split(',')) for c in links))
check('Revisión 1.1: guardas en llamadas de colaboración',all(any(guard in c.get('value') for c in links) for guard in ('[aprobada]','[rechazada]','[hay aprobadas]')))
_,stack=cells(stems[0])
check('Revisión 1.1: tres zonas stack delimitadas',all(stack[id].get('kind')=='zone' for id in ('planta','perimetro','nube')))
contain=True
for c in stack.values():
    if c.get('kind')!='component':continue
    g=c.find('mxGeometry'); z=stack[c.get('zone')].find('mxGeometry')
    x,y,w,h=[float(g.get(k)) for k in ('x','y','width','height')]; zx,zy,zw,zh=[float(z.get(k)) for k in ('x','y','width','height')]
    contain &= zx<=x and zy<=y and x+w<=zx+zw and y+h<=zy+zh
check('Revisión 1.1: componentes dentro de su zona',contain)
check('Revisión 1.1: router separado de Internet',stack['router'].get('zone')=='perimetro' and 'Router' not in stack['internet'].get('value'))
check('Revisión 1.1: HTTPS exclusivamente en enlace',any(c.get('edge')=='1' and c.get('protocol')=='HTTPS' for c in stack.values()) and not any('HTTPS' in c.get('value','') for c in stack.values() if c.get('kind')=='component'))
check('Revisión 1.1: nube condicionada y RPO/RTO propuestos',all(t in stack['note'].get('value') for t in ['condicionada','RPO 24 h','RTO 8 h','no garantías']))
baseline=json.loads((B/'evidencias/revision-1.1/linea-base.json').read_text(encoding='utf-8'))
base_hash={r['path']:r['hash'] for r in baseline['archivos']}
check('Revisión 1.1: tres factibilidades sin cambios',all(hashlib.sha256((B/p).read_bytes()).hexdigest().upper()==base_hash[p] for p in ['factibilidad/01-operativa.md','factibilidad/02-tecnica.md','factibilidad/03-economica.md']))
hash_original=hashlib.sha256((B/'enunciado/Simulacro de examen .pdf').read_bytes()).hexdigest().upper()
check('PDF original preservado y SHA-256 correcto',hash_original=='97B0B4F60A010FCA32062FBD7240A1505275625C6FA46A7FCAD14005F0F1C239')
check('Original: dos páginas completas',len(PdfReader(B/'enunciado/Simulacro de examen .pdf').pages)==2)
doc=B/'documentos/Solucion-Textiles-Globales.docx'
if doc.exists():
    dd=Document(doc)
    check('DOCX: seis diagramas insertados',len(dd.inline_shapes)==6)
    check('DOCX: trece apartados principales',sum(p.style.name=='Heading 1' for p in dd.paragraphs)==13)
    doc_text='\n'.join(p.text for p in dd.paragraphs)+'\n'+'\n'.join(cell.text for t in dd.tables for row in t.rows for cell in row.cells)
    check('Revisión 1.1: terminología corregida en DOCX',all(s in doc_text for s in ['Datos de corte y estado del registro','Datos de paquete habilitado y estado de confección','Liberar aprobadas para empaque','Código de lote y confirmación de liberación','P1 consulta D4']))
    with zipfile.ZipFile(doc) as zz:
        embedded={hashlib.sha256(zz.read(p)).hexdigest() for p in zz.namelist() if p.startswith('word/media/')}
    check('Revisión 1.1: seis láminas actuales insertadas, SHA idéntico',embedded=={hashlib.sha256((B/'diagramas/exportados'/f'{stem}.png').read_bytes()).hexdigest() for stem in stems})
final=B/'documentos/Solucion-Textiles-Globales.pdf'
if final.exists():
    rr=PdfReader(final)
    check('Documento PDF válido y no vacío',len(rr.pages)>10)
    check('Documento PDF sin páginas vacías',all(len((p.extract_text() or '').strip())>8 for p in rr.pages))
    pdf_text=norm(' '.join(p.extract_text() or '' for p in rr.pages))
    check('Revisión 1.1: terminología corregida en PDF final',all(s in pdf_text for s in ['Datos de corte y estado del registro','Datos de paquete habilitado y estado de confección','Liberar aprobadas para empaque','Código de lote y confirmación de liberación']))
patterns=[r'gh[pousr]_[A-Za-z0-9]{30,}',r'github_pat_[A-Za-z0-9_]{30,}',r'-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----',r'AKIA[0-9A-Z]{16}']
hits=[]
for p in B.rglob('*'):
    if p.suffix.lower() not in ('.md','.txt','.py','.ps1','.xml','.drawio','.json','.svg'):continue
    content=p.read_text(encoding='utf-8',errors='replace')
    if any(re.search(pattern,content) for pattern in patterns):hits.append(str(p.relative_to(B)))
check('Escaneo de credenciales: sin patrones reconocidos',not hits)
result={'sha256_original':hash_original,'alcance':'Sintaxis/estructura; requiere revisión académica y visual independiente','comprobaciones':checks,'fallos':sum(x['resultado']=='FAIL' for x in checks),'archivos_con_patrones':hits}
(B/'validacion/comprobaciones.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
raise SystemExit(1 if result['fallos'] else 0)
