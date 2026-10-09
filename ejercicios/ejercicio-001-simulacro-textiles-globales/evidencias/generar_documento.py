"""Informe editable: prosa Letter; láminas ampliadas para conservar legibilidad."""
from pathlib import Path
import re
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

B=Path(__file__).resolve().parents[1]
d=Document()
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Header','Footer']:
    st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',25),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
    d.styles[name].font.size=Pt(size)
    d.styles[name].paragraph_format.keep_with_next=True
def configure(sec,w=8.5,h=11,margin=.7):
    sec.page_width=Inches(w); sec.page_height=Inches(h)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(margin)
    sec.header_distance=sec.footer_distance=Inches(.3)
configure(d.sections[0])
d.sections[0].header.paragraphs[0].text='LABORATORIO ACADÉMICO  |  EJERCICIO 001'
f=d.sections[0].footer.paragraphs[0]
f.add_run('Textiles Globales S.A. · Versión final revisada  |  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
for p in (d.sections[0].header.paragraphs[0],f):
    for r in p.runs:r.font.size=Pt(9)

def clean(s):
    return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',s).replace('`','').replace('**','')
def table(rows):
    t=d.add_table(rows=1,cols=len(rows[0])); t.autofit=False
    width=(d.sections[-1].page_width-d.sections[-1].left_margin-d.sections[-1].right_margin)/914400
    n=len(rows[0]); weights=[1]*n
    if n==3:weights=[1.1,1.4,3.8]
    if n==3 and rows[0][0]=='ID':weights=[.55,1.9,4.55]
    if n==3 and rows[0][0]=='Fuente':weights=[2.8,1.5,2.7]
    if n==4:weights=[1.5,1.5,2.5,2.4]
    if n==5:weights=[1.9,2.5,2.5,1.8,2.3]
    if n==5 and rows[0][0]=='ID':weights=[.5,1.7,2.2,2.8,2.8]
    if n==6:weights=[1.4,1.2,1.2,1.7,1,1]
    for col,w in zip(t.columns,weights):col.width=Inches(width*w/sum(weights))
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,val) in enumerate(zip(cells,row)):
            cell.width=Inches(width*weights[j]/sum(weights)); cell.text=clean(val)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); margins.append(e)
            tcpr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DAE8FC'); tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
                p.paragraph_format.keep_with_next=i==0
                for run in p.runs:run.font.size=Pt(9 if n>4 else 10); run.bold=i==0
        trpr=t.rows[i]._tr.get_or_add_trPr()
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        trpr.append(OxmlElement('w:cantSplit'))
    d.add_paragraph().paragraph_format.space_after=Pt(2)

def md(path,skip_title=True):
    lines=(B/path).read_text(encoding='utf-8').splitlines(); i=0
    while i<len(lines):
        s=lines[i].strip(); i+=1
        if not s:continue
        if s.startswith('|'):
            rows=[]
            while True:
                if not re.fullmatch(r'[|\s:\-]+',s):rows.append([x.strip() for x in s.strip('|').split('|')])
                if i>=len(lines) or not lines[i].strip().startswith('|'):break
                s=lines[i].strip(); i+=1
            table(rows); continue
        if s.startswith('# '):
            if not skip_title:d.add_heading(clean(s[2:]),2)
        elif s.startswith('## '):d.add_heading(clean(s[3:]),2)
        elif s.startswith('### '):d.add_heading(clean(s[4:]),3)
        elif s.startswith('- '):d.add_paragraph(clean(s[2:]),style='List Bullet')
        else:d.add_paragraph(clean(s))

d.add_heading('1. Portada breve',1)
d.add_paragraph('Automatización del proceso de manufactura',style='Title')
d.add_paragraph('Textiles Globales S.A.\nSolución académica · Simulacro de examen, variante U',style='Subtitle')
d.add_paragraph('Universidad Mariano Gálvez de Guatemala · Seminario de Análisis y Diseño de Sistemas')
d.add_paragraph('Catedrático: Heriberto Antonio Escobar Menéndez\nFecha del examen original: 5 de octubre de 2026\nFecha de elaboración: 8 de octubre de 2026')
d.add_paragraph('Estudiante: ____________________\nCarné: ____________________')
d.add_paragraph('Fuente primaria: PDF original de dos páginas, conservado sin cambios. El texto recibido precisa la entrega, no reemplaza el enunciado.')
d.add_paragraph('Estado: versión final revisada del simulacro académico. No constituye una implementación ni una inversión autorizada.')
d.add_paragraph('Lectura: apartados 2–6, análisis y factibilidad; 7–12, seis láminas; 13, conclusiones, evidencia y pendientes. La prosa utiliza Letter; las láminas tienen formato ampliado proporcional para evitar reducir sus etiquetas.')
d.add_page_break()
d.add_heading('2. Caso',1); md('analisis/ANALISIS.md')
for number,title,path in [(3,'Factibilidad operativa','factibilidad/01-operativa.md'),(4,'Factibilidad técnica','factibilidad/02-tecnica.md'),(5,'Factibilidad económica','factibilidad/03-economica.md')]:
    d.add_heading(f'{number}. {title}',1); md(path)
d.add_heading('6. Conclusión de viabilidad',1)
d.add_paragraph('La viabilidad operativa exige capacitación por rol, captura intuitiva y un piloto acompañado. La viabilidad técnica exige equipos nuevos, cobertura medida, conectividad y respaldo probado. La opción recomendada es una aplicación web monolítica en nube administrada, condicionada a Internet y energía confiables; si no se verifican, debe reevaluarse una solución local.')
d.add_paragraph('La viabilidad económica no está demostrada con cifras. Obtener cotizaciones de I y O_anual, medir beneficios A y aprobar el umbral A ≥ O_anual + I/N antes de invertir. Disposición a financiar no equivale a rentabilidad. No se prometen reducciones porcentuales de merma ni de defectos.')

diagrams=[
('01-stack-tecnologias','Stack de tecnologías','Tres zonas: Planta, Perímetro y Nube. Router/firewall separado de Internet; HTTPS es el protocolo del enlace. Frontend y backend pertenecen al mismo servicio; PostgreSQL mantiene acceso privado. Respaldo y monitoreo cubren datos y aplicación. Nube condicionada a conectividad; RPO 24 h y RTO 8 h son objetivos propuestos.'),
('02-dfd-contexto','DFD de Contexto','Un proceso 0; cuatro roles externos al software; ocho flujos de datos nominales F01–F08. Las transferencias físicas de tela y prendas no se dibujan como flujos informáticos.'),
('03-dfd-nivel-0','DFD Nivel 0','Cuatro procesos y cuatro almacenes lógicos. Balance externo idéntico al contexto. P4 actualiza D4; P1 consulta D4 para elaborar F02 al Jefe. Se elimina P4→P1; la representación repetida de D4 abrevia D4→P1 sin agregar un almacén.'),
('04-casos-uso','Casos de Uso','Cinco dependencias include hacia comportamiento obligatorio. Registrar causa de rechazo extiende Registrar inspección solo cuando la prenda está rechazada. Liberar aprobadas para empaque incluye Generar código de lote y exige aprobadas; no atribuye empaque físico al Inspector ni inventa reproceso.'),
('05-secuencia','Secuencia','Escenario completo de producción. loop 13–17 por prenda; alt 14/15 mutuamente excluyentes. Mensajes 1–19 preservados; 18–20 solo si hay aprobadas. El 20, Código de lote y confirmación de liberación, confirma la liberación registrada, no el empaque físico. Sistema y Persistencia representan fronteras lógicas, no clases de implementación.'),
('06-colaboracion','Colaboración','Mismos participantes y 20 mensajes de secuencia. Los enlaces indican llamadas abreviadas numeradas y dirección; la tabla conserva los contenidos completos. Guardas aprobada, rechazada y hay aprobadas preservadas. Liberar aprobadas para empaque y confirmarLiberacion() no afirman ejecución de empaque físico.')]
for number,(stem,title,explanation) in enumerate(diagrams,7):
    im=Image.open(B/'diagramas/exportados'/f'{stem}.png')
    image_width=18.6; image_height=image_width*im.height/im.width
    sec=d.add_section(WD_SECTION_START.NEW_PAGE); configure(sec,20,image_height+2.5,.6)
    d.add_heading(f'{number}. {title}',1); d.add_paragraph(explanation)
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(B/'diagramas/exportados'/f'{stem}.png'),width=Inches(image_width))
    d.add_paragraph(f'Fuente oficial editable: diagramas/fuentes/{stem}.drawio · PDF individual: diagramas/exportados/{stem}.pdf')

configure(d.add_section(WD_SECTION_START.NEW_PAGE))
d.add_heading('13. Conclusiones',1)
d.add_paragraph('El modelo se limita a producción: habilitación, corte, confección, inspección y empaque. Los controles del PDF tienen trazabilidad en los seis artefactos; no se añadieron compras, ventas, RRHH ni requisitos mínimos DERCAS que este examen no impone.')
d.add_paragraph('La solución ofrece una base verificable para discutir captura, trazabilidad y medición. Renderizar un archivo no demuestra corrección académica: la aprobación de inferencias, tolerancias, fórmula de merma y semántica de lotes requiere revisión del docente o del responsable de planta.')
d.add_heading('Supuestos y ambigüedades para revisión',2); md('analisis/SUPOSICIONES.md')
d.add_paragraph('Otros pendientes humanos: completar nombre/carné; aprobar estaciones y turnos, cobertura, tarifas y horizonte de recuperación; acordar RPO/RTO; confirmar liberaciones parciales y política de rechazadas. La instrucción del PDF de agregar hescobarm@miumg.edu.gt como colaborador queda pendiente: no se compartió información fuera del trabajo local.')
d.add_heading('Evidencia de origen y diferencias',2); md('enunciado/PROCEDENCIA.md'); md('enunciado/COMPARACION.md')
d.add_heading('Consistencia cruzada',2); md('validacion/CONSISTENCIA.md')
d.add_heading('Fuentes y clasificación de evidencia',2); md('analisis/FUENTES.md')
d.add_heading('Validación y límites',2)
d.add_paragraph('Verificaciones reproducibles: validacion/comprobaciones.json y validacion/RESULTADOS.md. Incluyen XML DRAW.IO, referencias, lienzos, ocho flujos balanceados, direcciones include/extend, 20 mensajes equivalentes, exportaciones, hash original y patrones de credenciales. La revisión visual se realiza por separado sobre las seis exportaciones y todas las páginas del PDF final.')
d.add_paragraph('Entrega deliberadamente detenida antes de commit y publicación. Mantener rama codex/ejercicio-001-textiles-globales hasta revisión humana.')
d.core_properties.title='Solución académica — Textiles Globales S.A.'
d.core_properties.subject='Factibilidad, arquitectura, DFD y UML'
d.core_properties.author='Laboratorio académico'
d.core_properties.language='es-GT'
out=B/'documentos/Solucion-Textiles-Globales.docx'; out.parent.mkdir(exist_ok=True)
d.save(out)
print(out)
